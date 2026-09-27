"""The quarter's pack for the gestor: one ZIP, everything they ask for, nothing to type.

  Libros_<periodo>.xlsx   Resumen (303 / 130), Emitidas and Recibidas -- the two VAT
                          record books (libros registro), one row per invoice, with
                          totals that are formulas, so the gestor can check them.
  Facturas emitidas/      the PDF of every invoice issued in the quarter
  Gastos/                 the photo or PDF of every supplier bill that has one
  LEEME.txt               what is in it, and what the software cannot know

It can be downloaded from the panel, sent to the chat, or emailed straight to the
gestor -- the part of quarter-end an autónomo dreads most, done in one tap.
"""

import os
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Optional

from src import taxes
from src.config_loader import get_config

PACKS_DIR = Path("data/gestor")
# Gmail refuses attachments over 25 MB; leave room for the encoding overhead.
EMAIL_LIMIT = 18 * 1024 * 1024


def _safe(text: str) -> str:
    return re.sub(r"[^\w\-.]+", "_", text or "").strip("_") or "archivo"


def _state(record: dict) -> str:
    if record.get("rectified_by"):
        return f"Anulada por {record['rectified_by']}"
    if record["invoice"].rectifies:
        return "Rectificativa"
    return "Cobrada" if record.get("paid_at") else "Pendiente de cobro"


def build_workbook(year: int, quarter: int, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    from src.totals import breakdown

    config = get_config()
    start, end = taxes.quarter_bounds(year, quarter)
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1A3A5C")
    money = '#,##0.00 "€";-#,##0.00 "€"'

    def header(ws, titles, widths):
        ws.append(titles)
        for col, width in enumerate(widths, 1):
            cell = ws.cell(row=1, column=col)
            cell.font, cell.fill = head_font, head_fill
            cell.alignment = Alignment(vertical="center")
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = "A2"

    def totals_row(ws, label_col, money_cols):
        last = ws.max_row
        row = last + 1
        ws.cell(row=row, column=label_col, value="TOTAL").font = Font(bold=True)
        for col in money_cols:
            letter = get_column_letter(col)
            cell = ws.cell(row=row, column=col, value=f"=SUM({letter}2:{letter}{last})")
            cell.font, cell.number_format = Font(bold=True), money

    wb = Workbook()

    # ── Resumen ──────────────────────────────────────────────────────────────
    summary = wb.active
    summary.title = "Resumen"
    summary.column_dimensions["A"].width = 44
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 18
    rows = [
        [f"{config.name} — NIF {config.cif}"],
        [f"Periodo: {taxes.label(year, quarter)} ({start.strftime('%d/%m/%Y')} – "
         f"{end.strftime('%d/%m/%Y')})"],
        [f"Plazo de presentación: {taxes.deadline_text(year, quarter)}"],
        [f"Generado el {datetime.now().strftime('%d/%m/%Y %H:%M')}"],
        [],
        ["MODELO 303 — IVA", "Base", "Cuota"],
    ]
    vat = taxes.vat_return(year, quarter)
    for rate in sorted(vat.by_rate, reverse=True):
        if rate:
            base, tax = vat.by_rate[rate]
            rows.append([f"{taxes.rate_name(rate)} (devengado)", base, tax])
    for name, base in taxes.without_vat_rows(vat):
        rows.append([name, base, 0.0])
    rows.append(["Total IVA devengado", vat.output_base, vat.output_tax])
    rows.append(["IVA deducible (gastos registrados)", vat.input_base, vat.input_tax])
    rows.append(["Resultado (positivo = a ingresar)", None, vat.result])
    rows.append([])
    if taxes.is_company(config.cif):
        rows.append(["MODELO 130: no aplica a sociedades (pagos a cuenta: modelo 202)"])
    else:
        irpf = taxes.irpf_instalment(year, quarter)
        rows += [
            ["MODELO 130 — IRPF (acumulado desde el 1 de enero)", None, "Importe"],
            ["[01] Ingresos computables", None, irpf.income],
            ["[02] Gastos fiscalmente deducibles (registrados)", None, irpf.expenses],
            ["[03] Rendimiento neto", None, irpf.net],
            ["[04] 20 % del rendimiento neto", None, irpf.twenty_percent],
            ["[05] Pagos fraccionados de trimestres anteriores", None,
             irpf.previous_payments],
            ["[06] Retenciones e ingresos a cuenta soportados", None, irpf.withheld],
            ["[07] Resultado", None, irpf.result],
        ]
    rows += [
        [],
        ["Cálculo orientativo con las facturas y gastos registrados en el programa."],
        ["No incluye lo que no se ha registrado (cuota de autónomos, gastos bancarios…)."],
    ]
    if vat.bills_without_vat:
        rows.append([f"Atención: {vat.bills_without_vat} gasto(s) sin IVA desglosado."])
    for row in rows:
        summary.append(row)
    for row in summary.iter_rows(min_row=1, max_row=summary.max_row):
        for cell in row[1:]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = money
        if row[0].value and str(row[0].value).startswith("MODELO"):
            for cell in row:
                cell.font = Font(bold=True)
    summary["A1"].font = Font(bold=True, size=13)

    # ── Facturas emitidas ────────────────────────────────────────────────────
    issued = wb.create_sheet("Emitidas")
    header(issued, ["Fecha", "Número", "Cliente", "NIF cliente", "Base imponible",
                    "Tipo IVA %", "Cuota IVA", "Tipo IRPF %", "Retención IRPF",
                    "Total factura", "A cobrar", "Rectifica a", "Estado"],
           [11, 15, 32, 14, 15, 10, 13, 11, 14, 14, 14, 14, 20])
    records = taxes.issued_between(start, end)
    for record in records:
        inv = record["invoice"]
        t = breakdown(inv, config)
        issued.append([inv.date, inv.invoice_number, inv.client_name, inv.client_id or "",
                       t.base, t.tax_rate, t.tax, t.irpf_rate, t.irpf, t.gross, t.total,
                       inv.rectifies or "", _state(record)])
    for row in issued.iter_rows(min_row=2):
        row[0].number_format = "DD/MM/YYYY"
        for idx in (4, 6, 8, 9, 10):
            row[idx].number_format = money
    if records:
        totals_row(issued, 4, (5, 7, 9, 10, 11))

    # ── Facturas recibidas ───────────────────────────────────────────────────
    received = wb.create_sheet("Recibidas")
    header(received, ["Fecha", "Nº factura proveedor", "Proveedor", "NIF proveedor",
                      "Categoría", "Base imponible", "Cuota IVA", "Total", "Pagada",
                      "Documento"],
           [11, 20, 30, 14, 14, 15, 13, 13, 11, 40])
    bills = taxes.bills_between(start, end)
    for bill in bills:
        received.append([
            date.fromisoformat(bill["date"]), bill["reference"] or "",
            bill["supplier_name"], bill.get("supplier_tax_id") or "",
            bill["category"] or "", bill["subtotal"], bill["tax_amount"], bill["total"],
            "Sí" if bill["paid_at"] else "No",
            _bill_filename(bill) or "",
        ])
    for row in received.iter_rows(min_row=2):
        row[0].number_format = "DD/MM/YYYY"
        for idx in (5, 6, 7):
            row[idx].number_format = money
    if bills:
        totals_row(received, 5, (6, 7, 8))

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _bill_filename(bill: dict) -> Optional[str]:
    path = bill.get("file_path")
    if not path or not os.path.exists(path):
        return None
    return f"Gastos/{bill['date']}_{_safe(bill['supplier_name'])}_{bill['id']}" \
           f"{Path(path).suffix.lower()}"


def build_pack(year: int, quarter: int, out_dir: Optional[Path] = None) -> Path:
    """Build the ZIP for a quarter and return its path."""
    from src import finalize

    config = get_config()
    period = f"{year}-{quarter}T"
    out_dir = Path(out_dir or PACKS_DIR)
    out_dir.mkdir(parents=True, exist_ok=True)
    zip_path = out_dir / f"Gestor_{_safe(config.name)}_{period}.zip"
    workbook = build_workbook(year, quarter, out_dir / f"Libros_{period}.xlsx")

    start, end = taxes.quarter_bounds(year, quarter)
    records = taxes.issued_between(start, end)
    bills = taxes.bills_between(start, end)
    missing: list[str] = []

    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as pack:
        pack.write(workbook, workbook.name)
        for record in records:
            number = record["invoice"].invoice_number
            try:
                pdf = finalize.pdf_for(number)
            except Exception:
                pdf = None
            if pdf and os.path.exists(pdf):
                pack.write(pdf, f"Facturas emitidas/Factura_{_safe(number)}.pdf")
            else:
                missing.append(f"PDF de la factura {number}")
        for bill in bills:
            name = _bill_filename(bill)
            if name:
                pack.write(bill["file_path"], name)
            else:
                missing.append(f"documento del gasto de {bill['supplier_name']} "
                               f"({bill['date']}, {bill['total']:.2f} €)")
        pack.writestr("LEEME.txt", _readme(year, quarter, len(records), len(bills),
                                           missing))
    workbook.unlink(missing_ok=True)
    return zip_path


def _readme(year, quarter, n_invoices, n_bills, missing) -> str:
    config = get_config()
    lines = [
        f"Documentación del {taxes.label(year, quarter)} — {config.name} (NIF {config.cif})",
        f"Plazo de presentación: {taxes.deadline_text(year, quarter)}",
        "",
        f"Libros_{year}-{quarter}T.xlsx",
        "  Resumen: IVA (modelo 303) e IRPF (modelo 130), cálculo orientativo.",
        f"  Emitidas: libro registro de facturas expedidas ({n_invoices}).",
        f"  Recibidas: libro registro de facturas recibidas ({n_bills}).",
        "Facturas emitidas/  el PDF de cada factura del trimestre.",
        "Gastos/             la foto o el PDF de cada gasto que lo tiene.",
        "",
        "Solo incluye lo registrado en el programa. Añade lo que falte (cuota de",
        "autónomos, comisiones bancarias, gastos pagados sin factura registrada...).",
    ]
    if missing:
        lines += ["", "Sin documento adjunto:"] + [f"  - {m}" for m in missing]
    return "\r\n".join(lines) + "\r\n"


def email_to_gestor(year: int, quarter: int, to: Optional[str] = None) -> tuple[bool, str]:
    """Send the pack to the company's gestor. Returns (sent, what happened)."""
    from src.email_sender import send_email

    config = get_config()
    to = (to or getattr(config, "gestor_email", "") or "").strip()
    if not to:
        return False, ("No hay email del gestor configurado. Ponlo en el panel "
                       "(Configurar empresa → Email del gestor).")
    pack = build_pack(year, quarter)
    body = (f"Hola,\n\nTe envío la documentación del {taxes.label(year, quarter)} de "
            f"{config.name}.\n\n{taxes.summary_text(year, quarter)}\n\n"
            "En el ZIP van los libros registro en Excel, el PDF de cada factura "
            "emitida y los justificantes de los gastos.\n\nUn saludo,\n"
            f"{config.name}")
    subject = f"Documentación {taxes.label(year, quarter)} — {config.name}"
    note = ""
    attachment, name = pack, pack.name
    if pack.stat().st_size > EMAIL_LIMIT:
        # Too big for Gmail with every photo in it: send the books, keep the rest.
        attachment = build_workbook(year, quarter,
                                    pack.parent / f"Libros_{year}-{quarter}T.xlsx")
        name = attachment.name
        note = (" Solo van los libros en Excel: con todos los justificantes el ZIP "
                "pasaba del límite de Gmail. Descárgalo completo desde el panel.")
        body += "\n\n(Los justificantes no caben en un email: te los paso aparte.)"
    try:
        send_email(to=to, subject=subject, body=body, pdf_path=str(attachment),
                   attachment_name=name)
    except Exception as exc:
        return False, f"No se ha podido enviar el email: {exc}"
    _mark_sent(year, quarter, to)
    return True, f"Enviado a {to}.{note}"


def _mark_sent(year: int, quarter: int, to: str) -> None:
    from src import db

    with db.transaction() as conn:
        conn.execute("INSERT OR REPLACE INTO meta (key, value) VALUES (?, ?)",
                     (f"taxes:sent:{year}-{quarter}",
                      f"{datetime.now().isoformat(timespec='seconds')} {to}"))


def sent_on(year: int, quarter: int) -> Optional[str]:
    return _meta(f"taxes:sent:{year}-{quarter}")


def _meta(key: str) -> Optional[str]:
    from src import db

    row = db.connect().execute("SELECT value FROM meta WHERE key = ?", (key,)).fetchone()
    return row["value"] if row else None


def _set_meta(key: str) -> None:
    from src import db

    with db.transaction() as conn:
        conn.execute("INSERT OR REPLACE INTO meta (key, value) VALUES (?, ?)",
                     (key, datetime.now().isoformat(timespec="seconds")))


def buttons(year: int, quarter: int) -> list:
    """What can be done from the chat about a quarter."""
    row = [("📦 Paquete para el gestor", f"tax:pack:{year}-{quarter}")]
    if getattr(get_config(), "gestor_email", ""):
        row.append(("📧 Enviárselo al gestor", f"tax:mail:{year}-{quarter}"))
    return [row]


LAST_CALL_DAYS = 5


def quarter_reminder(today: Optional[date] = None) -> Optional[str]:
    """The daily check behind "the 303 is due": returns what it sent, if anything.

    On the first day of a filing window it sends the quarter's figures with a button
    for the pack. If, five days before the deadline, nothing has gone to the gestor
    yet, it says so once more. Each is sent once per quarter.
    """
    from src import telegram_access, telegram_api

    today = today or date.today()
    due = taxes.quarter_to_file(today)
    if due is None:
        return None
    year, quarter = due
    _, end = taxes.filing_window(year, quarter)

    first, last_call = f"taxes:reminded:{year}-{quarter}", f"taxes:lastcall:{year}-{quarter}"
    if not _meta(first):
        text = (f"📅 Ya se pueden presentar los impuestos del {taxes.label(year, quarter)}"
                f" ({taxes.deadline_text(year, quarter)}).\n\n"
                + taxes.summary_text(year, quarter))
        kind = "first"
        _set_meta(first)
    elif ((end - today).days <= LAST_CALL_DAYS and not sent_on(year, quarter)
          and not _meta(last_call)):
        days = (end - today).days
        text = (f"⏰ Quedan {days} día(s) para presentar el 303"
                f"{'' if taxes.is_company(get_config().cif) else ' y el 130'} del "
                f"{taxes.label(year, quarter)} y todavía no se le ha mandado nada al "
                "gestor.")
        kind = "last_call"
        _set_meta(last_call)
    else:
        return None

    for chat in telegram_access.notify_chats("taxes.view"):
        telegram_api.send_message(chat, text, buttons(year, quarter))
    return kind
