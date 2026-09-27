"""The quarter's taxes and the pack for the gestor.

The arithmetic is checked on invoices and bills with known answers, including the
awkward cases: a rectifying invoice cancelling one from the same quarter, two VAT
rates, an expense entered without its VAT, and the 130's running totals across the
year. Then the pack itself is opened and read, as the gestor would.
"""

import zipfile
from datetime import date

import pytest

from src import accounts, bills, finalize, gestor_pack, rectify, taxes
from src.config_loader import reload_config
from src.models import InvoiceData, InvoiceItem


def invoice_on(day: date, base: float, **kw):
    return InvoiceData(
        client_name=kw.pop("client_name", "Talleres Puig"),
        client_email=kw.pop("client_email", "taller@puig.es"),
        client_id=kw.pop("client_id", "B87654321"),
        items=[InvoiceItem("Servicio", 1, base)], date=day, **kw,
    )


@pytest.fixture
def autonomo():
    """A person (NIF, not CIF), so the 130 applies."""
    company = accounts.create_company("Laura Martí Disseny", tax_id="12345678Z")
    accounts.update_company(company, address="Carrer Nou 3, Igualada")
    reload_config()
    return company


@pytest.fixture
def sociedad():
    company = accounts.create_company("Talleres Mario S.L.", tax_id="B12345678")
    accounts.update_company(company, address="Pol. Les Comes 14, Igualada")
    reload_config()
    return company


# ── The calendar ─────────────────────────────────────────────────────────────

class TestCalendar:
    def test_quarter_bounds(self):
        assert taxes.quarter_bounds(2026, 3) == (date(2026, 7, 1), date(2026, 9, 30))
        assert taxes.quarter_bounds(2026, 4) == (date(2026, 10, 1), date(2026, 12, 31))

    def test_filing_windows(self):
        assert taxes.filing_window(2026, 3) == (date(2026, 10, 1), date(2026, 10, 20))
        assert taxes.filing_window(2026, 4) == (date(2027, 1, 1), date(2027, 1, 30))

    @pytest.mark.parametrize("today,expected", [
        (date(2026, 10, 5), (2026, 3)), (date(2026, 10, 20), (2026, 3)),
        (date(2026, 10, 21), None), (date(2026, 9, 27), None),
        (date(2027, 1, 29), (2026, 4)),
    ])
    def test_what_is_due_now(self, today, expected):
        assert taxes.quarter_to_file(today) == expected

    @pytest.mark.parametrize("tax_id,company", [
        ("B12345678", True), ("A08000000", True), ("12345678Z", False),
        ("X1234567L", False), ("", False),
    ])
    def test_company_or_person(self, tax_id, company):
        assert taxes.is_company(tax_id) is company


# ── Modelo 303 ───────────────────────────────────────────────────────────────

class TestVat:
    def test_output_by_rate_minus_deductible(self, offline, sociedad):
        finalize.issue(invoice_on(date(2026, 7, 10), 1000.0))             # 210
        finalize.issue(invoice_on(date(2026, 8, 3), 200.0, tax_rate=10))  # 20
        bills.create("Ferretería Puig", 121.0, subtotal=100.0,
                     bill_date=date(2026, 8, 20))                          # 21

        vat = taxes.vat_return(2026, 3)
        assert vat.by_rate == {21.0: [1000.0, 210.0], 10.0: [200.0, 20.0]}
        assert (vat.output_tax, vat.input_tax, vat.result) == (230.0, 21.0, 209.0)

    def test_other_quarters_are_left_out(self, offline, sociedad):
        finalize.issue(invoice_on(date(2026, 6, 30), 1000.0))
        finalize.issue(invoice_on(date(2026, 10, 1), 1000.0))
        assert taxes.vat_return(2026, 3).invoices == 0

    def test_a_cancelled_invoice_nets_to_zero(self, offline, sociedad):
        original = finalize.issue(invoice_on(date.today(), 500.0))
        rectify.rectify(original.number)
        year, quarter = taxes.quarter_of(date.today())
        assert taxes.vat_return(year, quarter).output_tax == 0

    def test_an_expense_without_its_vat_is_pointed_out(self, offline, sociedad):
        bills.create("Bar Pepe", 12.0, bill_date=date(2026, 7, 2))
        vat = taxes.vat_return(2026, 3)
        assert vat.bills_without_vat == 1
        assert "sin el IVA desglosado" in taxes.summary_text(2026, 3)


# ── Modelo 130 ───────────────────────────────────────────────────────────────

class TestIrpf:
    def test_running_totals_with_earlier_payments_and_withholding(self, offline, autonomo):
        finalize.issue(invoice_on(date(2026, 2, 10), 3000.0))               # Q1
        bills.create("Adobe", 60.5, subtotal=50.0, bill_date=date(2026, 2, 1))
        finalize.issue(invoice_on(date(2026, 5, 10), 2000.0, irpf_rate=15))  # Q2

        q1 = taxes.irpf_instalment(2026, 1)
        assert (q1.income, q1.expenses, q1.net) == (3000.0, 50.0, 2950.0)
        assert q1.to_pay == 590.0

        q2 = taxes.irpf_instalment(2026, 2)
        assert (q2.income, q2.net, q2.twenty_percent) == (5000.0, 4950.0, 990.0)
        assert (q2.previous_payments, q2.withheld) == (590.0, 300.0)
        assert q2.result == 100.0

    def test_a_loss_pays_nothing(self, offline, autonomo):
        bills.create("Ordenador", 1210.0, subtotal=1000.0, bill_date=date(2026, 1, 15))
        assert taxes.irpf_instalment(2026, 1).to_pay == 0

    def test_a_company_is_told_it_files_the_202_instead(self, offline, sociedad):
        text = taxes.summary_text(2026, 3)
        assert "modelo 202" in text and "Modelo 130 (IRPF" not in text

    def test_a_person_sees_the_130(self, offline, autonomo):
        assert "Modelo 130 (IRPF" in taxes.summary_text(2026, 3)


# ── The pack ─────────────────────────────────────────────────────────────────

class TestPack:
    @pytest.fixture
    def quarter(self, offline, sociedad, tmp_path, monkeypatch):
        monkeypatch.setattr(gestor_pack, "PACKS_DIR", tmp_path / "gestor")
        finalize.issue(invoice_on(date(2026, 7, 10), 1000.0))
        receipt = tmp_path / "ticket.jpg"
        receipt.write_bytes(b"\xff\xd8\xff fake jpeg")
        bills.create("Ferretería Puig", 121.0, subtotal=100.0,
                     bill_date=date(2026, 8, 20), file_path=str(receipt),
                     reference="F-88")
        bills.create("Bar Pepe", 12.0, bill_date=date(2026, 7, 2))
        return gestor_pack.build_pack(2026, 3)

    def test_it_holds_the_books_the_pdfs_and_the_receipts(self, quarter):
        names = zipfile.ZipFile(quarter).namelist()
        assert "Libros_2026-3T.xlsx" in names
        assert any(n.startswith("Facturas emitidas/Factura_") for n in names)
        assert any(n.startswith("Gastos/") and n.endswith(".jpg") for n in names)
        assert "LEEME.txt" in names

    def test_the_readme_lists_what_has_no_document(self, quarter):
        readme = zipfile.ZipFile(quarter).read("LEEME.txt").decode()
        assert "Bar Pepe" in readme

    def test_the_books_can_be_read_back(self, quarter, tmp_path):
        from openpyxl import load_workbook

        with zipfile.ZipFile(quarter) as pack:
            pack.extract("Libros_2026-3T.xlsx", tmp_path)
        wb = load_workbook(tmp_path / "Libros_2026-3T.xlsx")
        assert wb.sheetnames == ["Resumen", "Emitidas", "Recibidas"]

        issued = wb["Emitidas"]
        assert issued["C2"].value == "Talleres Puig"
        assert (issued["E2"].value, issued["G2"].value, issued["J2"].value) == \
            (1000.0, 210.0, 1210.0)
        assert issued["E3"].value == "=SUM(E2:E2)"  # totals are formulas

        received = wb["Recibidas"]
        assert received["C2"].value == "Bar Pepe"  # oldest first
        assert received["B3"].value == "F-88"


# ── Sending it ───────────────────────────────────────────────────────────────

class TestEmail:
    def test_without_a_gestor_email_it_says_where_to_put_one(self, offline, sociedad):
        sent, message = gestor_pack.email_to_gestor(2026, 3)
        assert not sent and "Email del gestor" in message

    def test_it_goes_with_the_zip_and_is_marked_sent(self, offline, sociedad, tmp_path,
                                                     monkeypatch):
        import src.email_sender as email_sender

        monkeypatch.setattr(gestor_pack, "PACKS_DIR", tmp_path / "gestor")
        accounts.update_company(sociedad, gestor_email="gestoria@ejemplo.es")
        reload_config()
        finalize.issue(invoice_on(date(2026, 7, 10), 1000.0))

        sent, message = gestor_pack.email_to_gestor(2026, 3)

        assert sent and "gestoria@ejemplo.es" in message
        to, subject, attachment = offline[-1]
        assert to == "gestoria@ejemplo.es" and "3T 2026" in subject
        assert attachment.endswith(".zip")
        assert gestor_pack.sent_on(2026, 3)

    def test_a_pack_too_big_for_gmail_sends_the_books(self, offline, sociedad, tmp_path,
                                                      monkeypatch):
        monkeypatch.setattr(gestor_pack, "PACKS_DIR", tmp_path / "gestor")
        monkeypatch.setattr(gestor_pack, "EMAIL_LIMIT", 10)
        accounts.update_company(sociedad, gestor_email="gestoria@ejemplo.es")
        reload_config()
        sent, message = gestor_pack.email_to_gestor(2026, 3)
        assert sent and "Solo van los libros" in message
        assert offline[-1][2].endswith(".xlsx")


# ── The reminder ─────────────────────────────────────────────────────────────

class TestReminder:
    @pytest.fixture(autouse=True)
    def owner(self, monkeypatch, sociedad):
        monkeypatch.setenv("TELEGRAM_CHAT_ID", "5001")
        reload_config()

    def sent_texts(self, outbox):
        return [p["text"] for m, p in outbox if m == "sendMessage"]

    def test_the_window_opening_is_announced_once(self, telegram_outbox):
        assert gestor_pack.quarter_reminder(date(2026, 10, 1)) == "first"
        assert gestor_pack.quarter_reminder(date(2026, 10, 2)) is None
        texts = self.sent_texts(telegram_outbox)
        assert len(texts) == 1 and "3T 2026" in texts[0]

    def test_a_last_call_if_nothing_went_to_the_gestor(self, telegram_outbox):
        gestor_pack.quarter_reminder(date(2026, 10, 1))
        assert gestor_pack.quarter_reminder(date(2026, 10, 16)) == "last_call"
        assert "Quedan 4 día(s)" in self.sent_texts(telegram_outbox)[-1]

    def test_no_last_call_once_it_was_sent(self, telegram_outbox):
        gestor_pack.quarter_reminder(date(2026, 10, 1))
        gestor_pack._mark_sent(2026, 3, "gestoria@ejemplo.es")
        assert gestor_pack.quarter_reminder(date(2026, 10, 16)) is None

    def test_outside_a_window_nothing_is_sent(self, telegram_outbox):
        assert gestor_pack.quarter_reminder(date(2026, 9, 27)) is None
        assert self.sent_texts(telegram_outbox) == []


# ── In the chat and on the web ───────────────────────────────────────────────

class TestSurfaces:
    def test_trimestre_in_the_chat_and_the_pack_button(self, offline, sociedad,
                                                       monkeypatch, tmp_path):
        from src import bot
        from test_telegram_access import press, say

        monkeypatch.setenv("TELEGRAM_CHAT_ID", "5001")
        reload_config()
        monkeypatch.setattr(gestor_pack, "PACKS_DIR", tmp_path / "gestor")
        finalize.issue(invoice_on(date(2026, 7, 10), 1000.0))

        chat = say(bot.cmd_trimestre, 5001, "2026", "3")
        assert "Modelo 303" in chat.text and "210,00" in chat.text
        assert "tax:pack:2026-3" in chat.buttons()

        pack = press(5001, "tax:pack:2026-3")
        assert pack.documents and pack.documents[0][0].endswith(".zip")

    def test_staff_without_the_permission_cannot_see_the_figures(self, sociedad,
                                                                 monkeypatch):
        from src import bot
        from test_telegram_access import linked_employee, say

        monkeypatch.setenv("TELEGRAM_CHAT_ID", "5001")
        reload_config()
        linked_employee(sociedad, ["invoices.view"])
        assert "No tienes permiso" in say(bot.cmd_trimestre, 7002).text

    def test_the_web_page_and_the_download(self, offline, sociedad, monkeypatch,
                                           tmp_path):
        import base64
        import json

        import itsdangerous
        from fastapi.testclient import TestClient

        monkeypatch.delenv("WEB_DEV_NO_AUTH", raising=False)
        monkeypatch.setenv("SESSION_SECRET", "test-secret")
        monkeypatch.setattr(gestor_pack, "PACKS_DIR", tmp_path / "gestor")
        from src.web import app as web_app

        accounts.create_user("jefe@talleres.es", accounts.ADMIN, company_id=sociedad)
        client = TestClient(web_app.app, follow_redirects=False)
        signer = itsdangerous.TimestampSigner("test-secret")
        data = base64.b64encode(json.dumps({"user": "jefe@talleres.es"}).encode())
        client.cookies.set("session", signer.sign(data).decode())
        finalize.issue(invoice_on(date(2026, 7, 10), 1000.0))

        page = client.get("/taxes?y=2026&q=3").text
        assert "Modelo 303" in page and "210,00" in page and "20/10/2026" in page

        download = client.get("/taxes/2026/3/pack")
        assert download.headers["content-type"] == "application/zip"
        assert zipfile.is_zipfile(__import__("io").BytesIO(download.content))
